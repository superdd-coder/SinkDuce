"""Opt-in conversion checks against an actual built application image.

Set SINKDUCE_OFFICE_TEST_IMAGE to run these without installing Office on the
test host. Only the disposable test directory and read-only source are mounted.
"""
import json
import os
from pathlib import Path
import subprocess

import pytest


pytestmark = pytest.mark.skipif(
    not os.environ.get("SINKDUCE_OFFICE_TEST_IMAGE"),
    reason="Set SINKDUCE_OFFICE_TEST_IMAGE for real Docker presentation conversion",
)


def _write_presentation(path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    for kind, title in [(XL_CHART_TYPE.COLUMN_CLUSTERED, "季度营收"),
                        (XL_CHART_TYPE.PIE, "市场份额")]:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
        for paragraph in slide.shapes.title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Noto Sans CJK SC"
        data = CategoryChartData()
        data.categories = ["Alpha", "Beta"]
        data.add_series("Revenue", (30, 70))
        chart = slide.shapes.add_chart(
            kind, Inches(1), Inches(1.5), Inches(8), Inches(5), data,
        ).chart
        chart.has_legend = False
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.show_value = True
        for point, color in zip(chart.series[0].points, [(31, 78, 121), (192, 0, 0)]):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor(*color)
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Table evidence"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(7), Inches(2)).table
    for row, values in zip(table.rows, [("Region", "Amount"), ("North", "12345")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    deck.save(path)


@pytest.fixture(scope="module")
def conversion(tmp_path_factory):
    work = tmp_path_factory.mktemp("office-runtime")
    _write_presentation(work / "charts.pptx")
    script = r'''
import base64, json, time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace
from src.chat_attachments import rendering, service
from src.chat_attachments.store import AttachmentStore

assert rendering.office_binary(), "Application image must include a presentation renderer"
work = Path("/verification")
store = AttachmentStore(Path("/tmp/sinkduce-office-test"))
source = (work / "charts.pptx").read_bytes()
attachment = store.create("slides", "charts.pptx", source)
start = time.monotonic()
service.parse_attachment(store, "slides", attachment.id, SimpleNamespace(enabled=True, api_token="unused"))
elapsed = time.monotonic() - start
item = store.get("slides", attachment.id)
assert item.status == "ready", item.error
refs = store.parsed("slides", item.id)["images"]
native = store.native_text("slides", item.id)
pdf = store.source_path("slides", item.id).parent / "pages.pdf"
assert pdf.is_file(), "PPTX must use whole-slide conversion rather than text-only fallback"
stamp = pdf.stat().st_mtime_ns
assert rendering.office_pdf(store.source_path("slides", item.id)) == pdf
assert pdf.stat().st_mtime_ns == stamp, "Reuse the cached converted PDF"
lazy = not list(pdf.parent.glob("image-*.txt"))
for index in range(len(refs)):
    url = store.image_url("slides", item.id, index)
    (work / f"slide-{index + 1}.jpg").write_bytes(base64.b64decode(url.split(",", 1)[1]))
unchanged = store.source_path("slides", item.id).read_bytes() == source

def convert(index):
    a = store.create("parallel", f"deck-{index}.pptx", source)
    service.parse_attachment(store, "parallel", a.id)
    result = store.get("parallel", a.id)
    return {"status": result.status, "warnings": result.warnings, "pages": result.image_count}

with ThreadPoolExecutor(max_workers=5) as executor:
    parallel = list(executor.map(convert, range(5)))
store.purge("slides")
store.purge("parallel")
print(json.dumps({"status": item.status, "warnings": item.warnings, "native": native,
    "refs": refs, "lazy": lazy, "source_unchanged": unchanged, "seconds": elapsed,
    "parallel": parallel, "purged": not any(store.root.iterdir())}, ensure_ascii=False))
'''
    repo = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        ["docker", "run", "--rm", "--network=none", "-i", "--entrypoint", "python",
         "--mount", f"type=bind,source={repo / 'src'},target=/app/src,readonly",
         "--mount", f"type=bind,source={work},target=/verification",
         "-e", "SINKDUCE_DATA=/tmp/sinkduce-test-data",
         os.environ["SINKDUCE_OFFICE_TEST_IMAGE"], "-"],
        input=script, capture_output=True, text=True, timeout=180,
    )
    payload = (json.loads(result.stdout.splitlines()[-1]) if result.returncode == 0
               else {"error": result.stdout + result.stderr})
    (work / "result.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2))
    return work, payload


def test_presentation_keeps_native_chinese_text_and_table_values(conversion):
    _, result = conversion
    assert "error" not in result, result.get("error")
    assert result["warnings"] == []
    for text in ["季度营收", "市场份额", "Alpha", "Beta", "30", "70", "North", "12345"]:
        assert text in result["native"]
    assert result["source_unchanged"]


@pytest.mark.parametrize("page", [1, 2])
def test_chart_data_is_visible_in_the_rendered_page(conversion, page):
    from PIL import Image

    work, result = conversion
    assert "error" not in result, result.get("error")
    with Image.open(work / f"slide-{page}.jpg") as image:
        width, height = image.size
        assert max(image.size) <= 2048
        plot = image.crop((width * .15, height * .25, width * .85, height * .8)).convert("RGB")
        rgb = plot.tobytes()
        pixels = list(zip(rgb[0::3], rgb[1::3], rgb[2::3]))
        # No legend is present: both colors must come from actual data marks.
        blue = sum(b > r + 25 and b > g + 5 for r, g, b in pixels)
        red = sum(r > g + 40 and r > b + 40 for r, g, b in pixels)
        assert blue > 1000, "Blue chart data is missing"
        assert red > 1000, "Red chart data is missing"


def test_all_slides_are_addressable_and_rasterization_stays_lazy(conversion):
    _, result = conversion
    assert "error" not in result, result.get("error")
    assert result["lazy"]
    assert [ref["slide_number"] for ref in result["refs"]] == [1, 2, 3]
    assert all(ref["render_source"] == "pages" for ref in result["refs"])


def test_five_presentations_convert_without_sharing_profiles(conversion):
    _, result = conversion
    assert "error" not in result, result.get("error")
    assert len(result["parallel"]) == 5
    assert all(item == {"status": "ready", "warnings": [], "pages": 3}
               for item in result["parallel"])
    assert result["purged"]
