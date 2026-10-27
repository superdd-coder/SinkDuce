"""Meeting downloads use the report document engine, with isolated fictional data."""
import io
from concurrent.futures import ThreadPoolExecutor
from threading import Condition, Event

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from src.meeting import store

DRAWING = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 320 100"><text x="12" y="40">Review → Release</text></svg>'
SUMMARY = ('## Summary\n\n[spk:0] approved the sample. [stt_0001]\n\n'
           '## Data & Facts\n\n| Item | Count |\n| --- | --- |\n| Sample | 17 |\n\n'
           '## Todo\n\n- [spk:0] Check labels [priority: high]\n\n'
           '## Detail\n\n```svg\n' + DRAWING + '\n```\n')


@pytest.fixture
def client(tmp_path, monkeypatch):
    from src.meeting.routes import router
    monkeypatch.setattr(store, 'MEETINGS_DIR', tmp_path / 'meetings')
    monkeypatch.setattr('src.identity.authorize', lambda *args, **kwargs: None)
    monkeypatch.setattr('src.identity.get_actor', lambda: None)
    meeting = store.create_meeting('Sample workshop')
    store.update_meeting(meeting.id, speaker_names={'0': 'Casey'}, tabs=[{'tab_id': 'tab_general', 'name': 'General', 'type': 'general'}])
    store.save_section_md(meeting.id, 'tab_general', SUMMARY)
    app = FastAPI(); app.include_router(router)
    with TestClient(app) as http:
        yield http, meeting


def export(http, meeting, fmt, **kwargs):
    return http.post(f'/meetings/{meeting.id}/sections/tab_general/export', json={'format': fmt, **kwargs})


def test_markdown_download_resolves_speakers_preserves_svg_and_does_not_modify_source(client):
    http, meeting = client
    response = export(http, meeting, 'md')
    assert response.status_code == 200
    assert response.headers['content-type'].startswith('text/markdown')
    assert 'attachment;' in response.headers['content-disposition']
    assert '.md' in response.headers['content-disposition']
    assert response.headers['access-control-expose-headers'] == 'Content-Disposition'
    assert 'Casey approved' in response.text
    assert '[stt_' not in response.text and '[priority:' not in response.text
    assert DRAWING in response.text and '| Sample | 17 |' in response.text
    assert store.get_section_md(meeting.id, 'tab_general') == SUMMARY


def test_word_download_contains_one_title_real_table_and_svg_picture(client):
    http, meeting = client
    response = export(http, meeting, 'docx')
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert '.docx' in response.headers['content-disposition']
    document = Document(io.BytesIO(response.content))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'Sample workshop' in text and 'Casey approved the sample.' in text
    assert document.tables[0].cell(1, 1).text == '17'
    assert len(document.inline_shapes) == 1
    assert 'Sources' not in text and '<svg' not in text
    assert '[stt_' not in text and '[priority:' not in text


def test_cleaning_cannot_expand_a_small_summary_past_export_limits(client):
    http, meeting = client
    from src.reports.models import MAX_MARKDOWN
    store.update_meeting(meeting.id, speaker_names={'0': 'x' * (MAX_MARKDOWN // 2 + 1)})
    store.save_section_md(meeting.id, 'tab_general', '[spk:0] [spk:0]')
    response = export(http, meeting, 'md')
    assert response.status_code == 422
    assert 'too large' in response.json()['detail']


def test_code_fences_remain_verbatim_and_long_minutes_have_no_automatic_report_contents(client):
    http, meeting = client
    code = '```svg\n' + DRAWING.replace('Review → Release', 'A  B [spk:0] [ref:7]') + '\n```'
    prose = '\n\n'.join('## Topic ' + str(i) + '\n\n' + ('Sample discussion. ' * 90) for i in range(5))
    store.save_section_md(meeting.id, 'tab_general', prose + '\n\n' + code)
    assert code in export(http, meeting, 'md').text
    document = Document(io.BytesIO(export(http, meeting, 'docx').content))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'Contents' not in text and 'Sources' not in text


def test_uploaded_images_are_embedded_and_other_files_are_not_accessible(client):
    http, meeting = client
    from PIL import Image
    image = io.BytesIO(); Image.new('RGB', (24, 12), 'navy').save(image, format='PNG')
    filename = 'sample.png'
    path = store.get_note_image_path(meeting.id, filename)
    path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(image.getvalue())
    url = f'/api/meetings/{meeting.id}/images/{filename}'
    for markup in [f'![Sample]({url})', f'<img src="{url}" alt="Sample">']:
        store.save_section_md(meeting.id, 'tab_general', markup)
        response = export(http, meeting, 'docx')
        assert response.status_code == 200
        assert len(Document(io.BytesIO(response.content)).inline_shapes) == 1
    for url in ['https://example.test/private.png', '/api/meetings/another/images/sample.png',
                f'/api/meetings/{meeting.id}/images/%2e%2e%2fmeta.json']:
        store.save_section_md(meeting.id, 'tab_general', f'![Sample]({url})')
        assert export(http, meeting, 'docx').status_code == 503


def test_pdf_download_uses_the_report_converter_and_surfaces_failures(client, monkeypatch):
    http, meeting = client
    seen = []
    def render(document, **kwargs):
        seen.append(document)
        return b'%PDF-1.7\nsynthetic download'
    monkeypatch.setattr('src.reports.rendering.render_pdf', render)
    response = export(http, meeting, 'pdf')
    assert response.status_code == 200 and response.content.startswith(b'%PDF-')
    assert response.headers['content-type'] == 'application/pdf'
    assert seen[0]['document_kind'] == 'meeting' and 'Casey approved' in seen[0]['markdown']
    def fail(*args, **kwargs):
        raise ValueError('PDF conversion unavailable')
    monkeypatch.setattr('src.reports.rendering.render_pdf', fail)
    failed = export(http, meeting, 'pdf')
    assert failed.status_code == 503 and 'unavailable' in failed.json()['detail']
    monkeypatch.setattr('src.reports.rendering.render_pdf', render)
    assert export(http, meeting, 'pdf').status_code == 200, 'A failed export releases its conversion slot'


def test_translation_export_uses_selected_saved_language_without_silent_fallback(client):
    http, meeting = client
    store.save_translation_md(meeting.id, 'tab_general', 'FR', '## Summary\n\nTranslated sample. [stt_0001]')
    response = export(http, meeting, 'md', language='fr')
    assert response.status_code == 200 and 'Translated sample.' in response.text
    assert 'Casey approved' not in response.text and 'FR' in response.headers['content-disposition']
    assert export(http, meeting, 'md', language='DE').status_code == 404


@pytest.mark.parametrize('body', [{'format': 'zip'}, {'format': 'md', 'language': '../../secret'}, {'format': 'pdf', 'language': ''}])
def test_unsupported_formats_and_unsafe_language_selectors_are_rejected(client, body):
    http, meeting = client
    assert http.post(f'/meetings/{meeting.id}/sections/tab_general/export', json=body).status_code == 422


def test_missing_meetings_or_sections_never_export_arbitrary_files(client):
    http, meeting = client
    assert http.post('/meetings/missing/sections/tab_general/export', json={'format': 'md'}).status_code == 404
    assert http.post(f'/meetings/{meeting.id}/sections/meta/export', json={'format': 'md'}).status_code == 404


def test_permission_is_checked_before_conversion(client, monkeypatch):
    http, meeting = client
    called = []
    monkeypatch.setattr('src.reports.rendering.render_pdf', lambda *a, **k: called.append(True))
    def deny(*args, **kwargs):
        raise HTTPException(403, 'forbidden')
    monkeypatch.setattr('src.identity.authorize', deny)
    assert export(http, meeting, 'pdf').status_code == 403
    assert called == []


def test_parallel_conversions_are_bounded_and_markdown_still_downloads(client, monkeypatch):
    http, meeting = client
    condition, release = Condition(), Event()
    started = []
    def slow(*args, **kwargs):
        with condition:
            started.append(True); condition.notify_all()
        assert release.wait(5)
        return b'%PDF-1.7\nsynthetic'
    monkeypatch.setattr('src.reports.rendering.render_pdf', slow)
    with ThreadPoolExecutor(max_workers=2) as pool:
        pending = [pool.submit(export, http, meeting, 'pdf') for _ in range(2)]
        try:
            with condition:
                assert condition.wait_for(lambda: len(started) == 2, timeout=3)
            assert export(http, meeting, 'pdf').status_code == 429
            assert export(http, meeting, 'md').status_code == 200
        finally:
            release.set()
        assert all(future.result().status_code == 200 for future in pending)
