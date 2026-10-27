import io
from pathlib import Path

import pytest
from docx import Document


def store(tmp_path, alive=None):
    from src.reports.store import ReportStore
    return ReportStore(tmp_path / "reports", session_exists=alive or (lambda sid: sid == "session-a"))


def test_report_cache_keeps_each_session_stores_liveness_check(tmp_path):
    from types import SimpleNamespace
    from src.reports.store import for_sessions

    first = SimpleNamespace(_db_path=tmp_path / "sessions.db", get_session=lambda sid: {"id": sid})
    second = SimpleNamespace(_db_path=tmp_path / "sessions.db", get_session=lambda sid: None)
    reports = for_sessions(first)
    reports.save("session-a", "Review", "Original.", [])
    reopened = for_sessions(second)

    assert reopened.list("session-a") == []
    assert len(reports.list("session-a")) == 1
    assert for_sessions(first) is reports
    assert for_sessions(second) is reopened


def test_report_versions_are_immutable_and_updates_use_optimistic_lock(tmp_path):
    s = store(tmp_path)
    first = s.save("session-a", "Review", "# First\n\nOriginal.", [], operation_id="call-1")
    second = s.save("session-a", "Review", "# Second\n\nRevised.", [], report_id=first["id"], expected_revision=1, operation_id="call-2")
    assert second["revision"] == 2
    assert s.read("session-a", first["id"], 1)["markdown"].endswith("Original.")
    assert s.read("session-a", first["id"], 2)["markdown"].endswith("Revised.")
    with pytest.raises(ValueError, match="version conflict"):
        s.save("session-a", "Review", "Stale edit", [], report_id=first["id"], expected_revision=1)


def test_tool_retry_is_idempotent_and_lists_one_report(tmp_path):
    s = store(tmp_path)
    a = s.save("session-a", "Review", "Body", [], operation_id="call-1")
    b = s.save("session-a", "Review", "Body", [], operation_id="call-1")
    assert a == b
    assert len(s.list("session-a")) == 1


def test_deleted_session_and_late_export_cannot_recreate_report(tmp_path):
    alive = {"session-a"}
    s = store(tmp_path, lambda sid: sid in alive)
    report = s.save("session-a", "Review", "Body", [])
    alive.clear()
    s.purge("session-a")
    with pytest.raises(FileNotFoundError):
        s.save("session-a", "Late report", "Body", [])
    with pytest.raises(FileNotFoundError):
        s.write_export("session-a", report["id"], 1, "html", b"late", status="ready")
    assert not (s.root / "session-a").exists()


def test_report_paths_and_unsupported_formats_are_rejected(tmp_path):
    s = store(tmp_path)
    with pytest.raises(ValueError):
        s.list("../outside")
    a = s.save("session-a", "Report", "Body", [])
    with pytest.raises(ValueError):
        s.read("session-a", "../../outside")
    with pytest.raises(ValueError):
        s.write_export("session-a", a["id"], 1, "../../outside", b"x", status="ready")


def test_export_state_is_version_and_format_specific_and_recovers_interruption(tmp_path):
    s = store(tmp_path)
    a = s.save("session-a", "Review", "Body", [])
    s.write_export("session-a", a["id"], 1, "html", b"ready", status="ready")
    s.write_export("session-a", a["id"], 1, "pdf", None, status="running")
    reopened = store(tmp_path)
    assert reopened.export_state("session-a", a["id"], 1, "html")["status"] == "ready"
    assert reopened.export_state("session-a", a["id"], 1, "pdf")["status"] == "failed"
    assert reopened.export_path("session-a", a["id"], 1, "html").read_bytes() == b"ready"


def test_export_from_an_old_renderer_is_retryable_instead_of_serving_a_stale_file(tmp_path):
    import json
    s = store(tmp_path)
    report = s.save("session-a", "Review", "Body", [])
    s.write_export("session-a", report["id"], 1, "html", b"old", status="ready")
    state_path = s.export_path("session-a", report["id"], 1, "html").with_suffix(".html.json")
    state = json.loads(state_path.read_text())
    state["renderer_version"] = "previous"
    state_path.write_text(json.dumps(state))
    assert s.export_state("session-a", report["id"], 1, "html")["status"] == "failed"


def test_exports_share_chinese_tables_and_real_citation_appendix(tmp_path):
    from src.chatbox.citations import EvidenceRegistry
    from src.reports.rendering import render_html, render_docx
    r = EvidenceRegistry("session-a")
    e = r.register("原文数值：基线 6 天，试点 4.2 天。", {"file_id": "file-a", "source_label": "试点.pdf", "page_number": 4})
    body = f"## 结果\n\n等待时间减少 30%。[cite:{e['id']}]\n\n| 阶段 | 天数 |\n|---|---|\n|基线|6|\n|试点|4.2|"
    report = {"title": "试点复盘", "markdown": body, "citations": [e], "revision": 1}
    html = render_html(report).decode()
    assert "试点复盘" in html and "<table" in html
    assert e["quote"] in html and "试点.pdf" in html
    assert "[cite:" not in html
    assert "source-" in html
    document = Document(io.BytesIO(render_docx(report)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "试点复盘" in text and e["quote"] in text
    assert "[cite:" not in text
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "基线"


def test_rendering_rejects_unknown_citations_and_strips_active_html():
    from src.reports.rendering import render_html
    with pytest.raises(ValueError, match="Unknown citation"):
        render_html({"title": "Report", "markdown": "Fake[cite:invented]", "citations": []})
    html = render_html({"title": "<script>bad</script>", "markdown": '<script>alert(1)</script><img src="file:///etc/passwd" onerror="alert(1)"><iframe src="https://invalid.test"></iframe>\n\nSafe.', "citations": []}).decode()
    assert "<script" not in html and "onerror=" not in html
    assert "file:///" not in html and "<iframe" not in html
    assert "Safe." in html
    with pytest.raises(ValueError, match="image references"):
        render_html({"title": "Report", "markdown": "![External](https://invalid.test/tracker.png)", "citations": []})


def test_safe_svg_is_embedded_in_html_and_word_but_remote_svg_is_rejected():
    from src.reports.rendering import render_html, render_docx
    report = {"title": "Diagram", "markdown": '```svg\n<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50"><rect width="50" height="40" fill="green"/></svg>\n```', "citations": []}
    assert "data:image/png;base64," in render_html(report).decode()
    assert len(Document(io.BytesIO(render_docx(report))).inline_shapes) == 1
    report["markdown"] = '```svg\n<svg><image href="https://invalid.test/x.png"/></svg>\n```'
    with pytest.raises(ValueError):
        render_html(report)


def test_word_citations_link_to_bookmarked_source_excerpts():
    import zipfile
    from src.chatbox.citations import EvidenceRegistry
    from src.reports.rendering import render_docx
    evidence = EvidenceRegistry("s").register("Original evidence", {"file_id": "a"})
    report = {"title": "Review", "markdown": f"Fact[cite:{evidence['id']}]", "citations": [evidence]}
    with zipfile.ZipFile(io.BytesIO(render_docx(report))) as archive:
        xml = archive.read("word/document.xml").decode()
    assert 'w:hyperlink' in xml and 'w:anchor="source_' in xml
    assert 'w:bookmarkStart' in xml


def test_long_reports_have_a_navigable_contents_section_in_both_exports():
    import zipfile
    from src.reports.rendering import render_html, render_docx
    body = "\n\n".join(f"## Section {index}\n\n" + "A detailed discussion. " * 100 for index in range(1, 5))
    report = {"title": "Long report", "markdown": body, "citations": []}
    html = render_html(report).decode()
    assert 'href="#section-1"' in html and 'id="section-1"' in html
    with zipfile.ZipFile(io.BytesIO(render_docx(report))) as archive:
        xml = archive.read("word/document.xml").decode()
    assert 'w:anchor="section_1"' in xml and 'w:name="section_1"' in xml


def test_report_tool_keeps_one_canonical_title_instead_of_duplicating_it(tmp_path):
    import json
    from src.chatbox.citations import EvidenceRegistry
    from src.db.sessions import SessionStore
    from src.reports.tools import execute_report_tool
    from src.reports.store import for_sessions
    sessions = SessionStore(str(tmp_path / "sessions.db"))
    sessions.create_session(session_id="session-a")
    result = json.loads(execute_report_tool(sessions, "session-a", "create_report",
        {"title": "Review", "markdown": "# Review\n\n## Findings\n\nBody."}, EvidenceRegistry("session-a")))
    record = for_sessions(sessions).read("session-a", result["report"]["id"])
    assert record["markdown"].startswith("## Findings")


@pytest.mark.parametrize("args", [
    {"title": "Review", "markdown": "x" * 120001},
    {"title": "x" * 201, "markdown": "Body"},
    {"title": "Review", "markdown": ["Body"]},
    {"title": "Review", "markdown": "   "},
])
def test_invalid_report_inputs_are_rejected_before_rendering(tmp_path, monkeypatch, args):
    import json
    from src.chatbox.citations import EvidenceRegistry
    from src.db.sessions import SessionStore
    from src.reports import tools
    sessions = SessionStore(str(tmp_path / "sessions.db"))
    sessions.create_session(session_id="session-a")
    def unexpected_render(_record):
        pytest.fail("Invalid input must not invoke document/SVG rendering")
    monkeypatch.setattr(tools, "render_html", unexpected_render)
    result = json.loads(tools.execute_report_tool(sessions, "session-a", "create_report", args, EvidenceRegistry("session-a")))
    assert "error" in result


def test_source_appendix_renders_managed_images_and_tables_in_html_and_word():
    from bs4 import BeautifulSoup
    from PIL import Image
    from src.chatbox.citations import EvidenceRegistry
    from src.reports.rendering import render_html, render_docx
    output = io.BytesIO()
    Image.new("RGB", (400, 160), "#1a5e3d").save(output, format="PNG")
    quote = ":::image\nimage_id: img-a\nfile_id: file-a\n:::\n\n| Item | Count |\n|---|---|\n| Sample | 42 |"
    e = EvidenceRegistry("s").register(quote, {"file_id": "file-a", "collection": "col"})
    report = {"title": "Review", "markdown": f"Finding[cite:{e['id']}]", "citations": [e]}
    resolver = lambda evidence, fields: output.getvalue()
    html = render_html(report, source_images=resolver).decode()
    assert ":::image" not in html and "image_id:" not in html
    soup = BeautifulSoup(html, "html.parser")
    assert soup.select_one(".source img")["src"].startswith("data:image/png;base64,")
    assert soup.select_one(".source table").get_text().find("42") >= 0
    doc = Document(io.BytesIO(render_docx(report, source_images=resolver)))
    assert len(doc.inline_shapes) == 1 and len(doc.tables) == 1
    assert not any(":::image" in p.text for p in doc.paragraphs)
    missing = render_html(report).decode()
    assert ":::image" not in missing and "image_id:" not in missing
    assert "图片暂不可用" in missing


def test_wide_word_tables_use_landscape_and_weighted_columns_then_restore_portrait():
    from docx.enum.section import WD_ORIENT
    from src.reports.rendering import render_docx
    # Entirely fictional inventory data; never copy uploaded business records.
    headers = ["No.", "Item", "Category", "Units", "Weight (kg)", "Budget", "Batch", "Ordered", "Received", "Checked", "Status"]
    row = ["1", "Crate-A", "Synthetic sample", "18", "42", "560", "LOT-A", "Q1-2034", "Q2-2034", "Q4-2034", "Awaiting quality inspection"]
    body = "## Sample inventory\n\n|" + "|".join(headers) + "|\n|" + "|".join(["---"] * 11) + "|\n|" + "|".join(row) + "|\n\n## Conclusion\n\nBody."
    doc = Document(io.BytesIO(render_docx({"title": "Review", "markdown": body, "citations": []})))
    assert any(s.orientation == WD_ORIENT.LANDSCAPE for s in doc.sections)
    assert doc.sections[-1].orientation == WD_ORIENT.PORTRAIT
    widths = [col.width for col in doc.tables[0].columns]
    assert widths[1] > widths[0] and widths[-1] > widths[0]
    assert 'w:tblHeader' in doc.tables[0]._tbl.xml


def test_incomplete_source_image_fence_does_not_leak_internal_fields_or_external_images():
    from src.chatbox.citations import EvidenceRegistry
    from src.reports.rendering import render_html
    e = EvidenceRegistry('s').register(':::image\nimage_id: missing\nfile_id: file-a', {'file_id':'file-a'})
    report = {'title':'Review', 'markdown':f"Fact[cite:{e['id']}]", 'citations':[e]}
    output = render_html(report).decode()
    assert ':::image' not in output and 'image_id:' not in output
    assert '图片暂不可用' in output
    e['quote'] = '![remote](https://invalid.test/tracker.png)<img src="file:///etc/passwd">'
    output = render_html(report).decode()
    assert '<img' not in output and 'file:///' not in output and 'tracker.png' not in output
def test_source_caption_page_must_match_its_bound_evidence(tmp_path):
    from src.reports.store import ReportStore
    from src.chatbox.citations import EvidenceRegistry
    store=ReportStore(tmp_path,session_exists=lambda sid:True)
    e=EvidenceRegistry('s').register('Verified source text.',{'file_id':'a','page_number':4})
    with pytest.raises(ValueError,match='[Cc]aption'):
        store.save('s','Report',f"来源：第 7 页结构图 [cite:{e['id']}]。",[e])
    assert store.save('s','Report',f"来源：第 4 页结构图 [cite:{e['id']}]。",[e])['revision']==1


def test_unmodified_legacy_caption_does_not_block_an_unrelated_local_edit(tmp_path):
    from src.reports.store import ReportStore
    from src.chatbox.citations import EvidenceRegistry
    store = ReportStore(tmp_path, session_exists=lambda _: True)
    evidence = EvidenceRegistry('s').register('Verified source.', {'file_id': 'a', 'page_number': 4})
    caption = f"来源：第 7 页结构图 [cite:{evidence['id']}]。"
    old = store.save('s', 'Report', 'Long paragraph.\n\n' + caption, [evidence], historical_replay=True)
    block = old['blocks'][0]
    new = store.save_edits('s', old['id'], [{'id':'shorten', 'block_id':block['id'], 'block_hash':block['hash'], 'before':'Long paragraph.', 'after':'Brief.'}], [evidence], expected_revision=1, operation_id='edit')
    assert new['markdown'] == 'Brief.\n\n' + caption
    # Altering the bad caption, duplicating it, or rebinding evidence remains invalid.
    for body, records in [
        (new['markdown'].replace('结构图', '新结构图'), [evidence]),
        (new['markdown'] + '\n\n' + caption, [evidence]),
        (new['markdown'], [{**evidence, 'location':{'page_number':2}}]),
    ]:
        with pytest.raises(ValueError, match='[Cc]aption|Unknown citation'):
            store.save('s', 'Report', body, records, report_id=old['id'], expected_revision=2)
    assert store.read('s', old['id'])['revision'] == 2



def test_export_uses_one_title_and_distinct_source_captions():
    from bs4 import BeautifulSoup
    from src.reports.rendering import render_html, render_docx
    report = {"title": "Research report", "markdown": "# Research draft\n\n**A concise introduction**\n\n## Findings\n\n**Funding details**\n\nBody.\n\n来源：原始资料。", "citations": []}
    soup = BeautifulSoup(render_html(report), "html.parser")
    assert [h.get_text() for h in soup.find_all("h1")] == ["Research report"]
    assert soup.select_one("p.report-caption").get_text() == "来源：原始资料。"
    assert soup.select_one("p.report-deck")
    doc = Document(io.BytesIO(render_docx(report)))
    assert "Research draft" not in [p.text for p in doc.paragraphs]
    caption = next(p for p in doc.paragraphs if p.text == "来源：原始资料。")
    assert caption.runs[0].font.size.pt < doc.styles["Normal"].font.size.pt
