import asyncio
import io
import json
import threading
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from PIL import Image


@pytest.fixture
def api(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from src.api.routes import sessions, chat_attachments
    from src.db.sessions import SessionStore
    from src.services import services
    s = SessionStore(str(tmp_path / "sessions.db"))
    s.create_session(session_id="owned", created_by="local")
    s.create_session(session_id="foreign", created_by="another-user")
    monkeypatch.setattr(services, "session_store", s)
    monkeypatch.setattr(services, "config", SimpleNamespace(mineru=None))
    monkeypatch.setattr("src.chatbox.image_memory.default_base_dir", lambda: tmp_path / "session_images")
    app = FastAPI()
    app.include_router(chat_attachments.router, prefix="/api")
    app.include_router(sessions.router, prefix="/api")
    return TestClient(app), s


def test_citation_text_preview_reads_existing_parsed_content_and_checks_session_access(api):
    from src.chat_attachments.store import for_sessions
    from src.identity.actor import Actor
    from unittest.mock import patch
    client, sessions = api
    a = for_sessions(sessions).create("owned", "Source.txt", b"Original fact.")
    for_sessions(sessions).finish("owned", a.id, content="Original fact.", position_map=[], images=[], parser="local")
    response = client.get(f"/api/sessions/owned/attachments/{a.id}/text")
    assert response.status_code == 200
    assert response.json()["text"] == "Original fact."
    assert client.get(f"/api/sessions/foreign/attachments/{a.id}/text").status_code == 404
    with patch("src.identity.get_actor", return_value=Actor(id="stranger", kind="key", role_id="member")), patch("src.identity.may_see_everything", return_value=False):
        assert client.get(f"/api/sessions/owned/attachments/{a.id}/text").status_code == 404


def test_api_upload_parse_download_and_explicit_delete(api):
    from src.chat_attachments.store import for_sessions
    client, sessions = api
    store = for_sessions(sessions)
    res = client.post("/api/sessions/owned/attachments", files={"file": ("facts.txt", b"context fact")})
    assert res.status_code == 201, res.text
    aid = res.json()["id"]
    rows = client.get("/api/sessions/owned/attachments").json()
    assert rows[0]["status"] == "ready"
    assert client.get(f"/api/sessions/owned/attachments/{aid}/source").content == b"context fact"
    assert (store.root / "owned" / aid / "parsed.md").is_file()
    assert client.delete("/api/sessions/owned").status_code == 204
    assert not (store.root / "owned").exists()
    assert client.get(f"/api/sessions/owned/attachments/{aid}/source").status_code == 404


@pytest.mark.parametrize("extension", ["ppt", "pptx"])
def test_presentation_preview_reuses_converted_pdf_and_download_keeps_source(api, extension):
    from src.chat_attachments.store import for_sessions
    client, sessions = api
    store = for_sessions(sessions)
    original = b"original presentation bytes"
    a = store.create("owned", f"deck.{extension}", original)
    store.finish("owned", a.id, content="Slide content", position_map=[],
                 images=[{"render_source": "pages", "page_number": 1}], parser="local")
    pdf = b"%PDF-1.4\nconverted presentation"
    (store.source_path("owned", a.id).parent / "pages.pdf").write_bytes(pdf)
    preview = client.get(f"/api/sessions/owned/attachments/{a.id}/preview")
    assert preview.status_code == 200, preview.text
    assert preview.headers["content-type"] == "application/pdf"
    assert preview.headers["content-disposition"].startswith("inline;")
    assert preview.content == pdf
    assert client.get(f"/api/sessions/owned/attachments/{a.id}/source").content == original


def test_presentation_without_converted_pdf_can_preview_its_original(api):
    from src.chat_attachments.store import for_sessions
    client, sessions = api
    store = for_sessions(sessions)
    a = store.create("owned", "deck.pptx", b"original presentation bytes")
    store.finish("owned", a.id, content="Slide content", position_map=[], images=[], parser="local")
    preview = client.get(f"/api/sessions/owned/attachments/{a.id}/preview")
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert preview.headers["content-disposition"].startswith("inline;")
    assert preview.content == b"original presentation bytes"


def test_five_attachment_parses_run_concurrently_and_a_sixth_waits(tmp_path, monkeypatch):
    from src.api.routes import chat_attachments

    release = threading.Event()
    five_started = threading.Event()
    lock = threading.Lock()
    active = started = peak = 0

    def parse(*args):
        nonlocal active, started, peak
        with lock:
            active += 1
            started += 1
            peak = max(peak, active)
            if active == 5:
                five_started.set()
        try:
            assert release.wait(10), "The test releases every parser even if concurrency is incorrect"
        finally:
            with lock:
                active -= 1

    monkeypatch.setattr(chat_attachments, "parse_attachment", parse)

    async def run():
        from src.chat_attachments.store import AttachmentStore
        store = AttachmentStore(tmp_path / "attachments")
        files = [store.create("session", f"sample-{i}.txt", b"fictional") for i in range(6)]
        jobs = [asyncio.create_task(chat_attachments._parse(store, "session", a.id, None)) for a in files]
        try:
            assert await asyncio.to_thread(five_started.wait, 2), "Five files should begin parsing without waiting for earlier files to finish"
            with lock:
                assert started == peak == 5, "The sixth file must remain queued"
        finally:
            release.set()
            await asyncio.gather(*jobs)
        assert started == 6
        assert peak == 5

    asyncio.run(run())


def test_api_denies_foreign_session_and_invalid_file(api):
    client, _ = api
    # Use explicit caller scope; local owner/admin is intentionally allowed all sessions.
    from src.identity.actor import Actor
    from unittest.mock import patch
    actor = Actor(id="unrelated", kind="key", role_id="member")
    with patch("src.identity.get_actor", return_value=actor), patch("src.identity.may_see_everything", return_value=False):
        assert client.get("/api/sessions/foreign/attachments").status_code == 404
    res = client.post("/api/sessions/owned/attachments", files={"file": ("binary.exe", b"bad")})
    assert res.status_code == 400


@pytest.mark.parametrize("filename,data", [
    ("notes.txt", b"Attachment fact"), ("notes.md", b"# Attachment fact"),
    ("data.csv", b"name,value\nAttachment,fact\n"), ("data.tsv", b"name\tvalue\nAttachment\tfact\n"),
    ("page.html", b"<h1>Attachment fact</h1>"), ("page.htm", b"<p>Attachment fact</p>"),
    ("data.json", b'{"Attachment":"fact"}'), ("data.jsonl", b'{"Attachment":"fact"}\n'),
])
def test_api_accepts_all_text_parser_types_without_collection_permissions(api, monkeypatch, filename, data):
    from src.services import services
    from src.chat_attachments.store import for_sessions
    db = MagicMock()
    db.get_collection_config.side_effect = AssertionError("Chat uploads must not query collection restrictions")
    monkeypatch.setattr(services, "db", db)
    client, sessions = api
    res = client.post("/api/sessions/owned/attachments", files={"file": (filename, data)})
    assert res.status_code == 201
    aid = res.json()["id"]
    assert for_sessions(sessions).get("owned", aid).status == "ready"
    assert "fact" in for_sessions(sessions).text("owned", aid)
    db.get_collection_config.assert_not_called()


@pytest.mark.parametrize("extension", ["docx", "xlsx", "pptx"])
def test_api_parses_real_office_files(api, extension):
    buf = io.BytesIO()
    if extension == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph("Attachment fact")
        doc.save(buf)
    elif extension == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["Attachment", "fact"])
        wb.save(buf)
        wb.close()
    else:
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Attachment fact"
        presentation.save(buf)
    client, sessions = api
    res = client.post("/api/sessions/owned/attachments", files={"file": (f"office.{extension}", buf.getvalue())})
    from src.chat_attachments.store import for_sessions
    a = for_sessions(sessions).get("owned", res.json()["id"])
    assert a.status == "ready", a.error
    assert "fact" in for_sessions(sessions).text("owned", a.id)


def test_api_parses_a_real_pdf_and_keeps_page_positions(api):
    stream = b"BT /F1 12 Tf 36 720 Td (Attachment fact) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    pdf = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(pdf)
    pdf += b"xref\n0 6\n0000000000 65535 f \n"
    pdf += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    pdf += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    client, sessions = api
    res = client.post("/api/sessions/owned/attachments", files={"file": ("source.pdf", pdf)})
    from src.chat_attachments.store import for_sessions
    store = for_sessions(sessions)
    a = store.get("owned", res.json()["id"])
    assert a.status == "ready", a.error
    assert "Attachment fact" in store.text("owned", a.id)
    assert store.parsed("owned", a.id)["position_map"][0]["page_number"] == 1


def test_api_rejects_size_limit_and_retries_failed_parse(api, monkeypatch):
    from src.api.routes import chat_attachments
    from src.chat_attachments import service
    from src.chat_attachments.store import for_sessions
    client, sessions = api
    monkeypatch.setattr(chat_attachments, "MAX_FILE_BYTES", 3)
    assert client.post("/api/sessions/owned/attachments", files={"file": ("a.txt", b"1234")}).status_code == 413
    res = client.post("/api/sessions/owned/attachments", files={"file": ("broken.pdf", b"bad")})
    aid = res.json()["id"]
    assert for_sessions(sessions).get("owned", aid).status == "failed"
    monkeypatch.setattr("src.chat_attachments.rendering.inspect_pdf", lambda path: ("Recovered text", [], []))
    assert client.post(f"/api/sessions/owned/attachments/{aid}/retry").status_code == 200
    assert for_sessions(sessions).get("owned", aid).status == "ready"
    for_sessions(sessions).bind("owned", [aid])
    assert client.delete(f"/api/sessions/owned/attachments/{aid}").status_code == 409


def test_image_preview_uses_bounded_rendition_and_original_download_is_lossless(api, monkeypatch):
    monkeypatch.setattr("src.parsers.image_utils._ocr_image", lambda *a, **kw: ("", 0))
    buf = io.BytesIO()
    Image.new("RGB", (3200, 2000), "red").save(buf, format="PNG")
    client, _ = api
    a = client.post("/api/sessions/owned/attachments", files={"file": ("photo.png", buf.getvalue())}).json()
    preview = client.get(f'/api/sessions/owned/attachments/{a["id"]}/preview')
    assert preview.status_code == 200
    with Image.open(io.BytesIO(preview.content)) as image:
        assert max(image.size) <= 2048
    assert client.get(f'/api/sessions/owned/attachments/{a["id"]}/source').content == buf.getvalue()


def test_attachment_only_message_reaches_agent_and_is_bound(api, monkeypatch):
    from src.services import services
    client, sessions = api
    a = client.post("/api/sessions/owned/attachments", files={"file": ("facts.txt", b"fact")}).json()
    calls = []
    class Agent:
        async def chat_stream(self, sid, content, **kw):
            calls.append((content, kw))
            yield {"type": "done"}
    monkeypatch.setattr(services, "chatbox_agent", Agent())
    res = client.post("/api/sessions/owned/messages", json={"content": "", "attachment_ids": [a["id"]]})
    assert res.status_code == 200, res.text
    assert calls[0][1]["attachment_ids"] == [a["id"]]
    assert calls[0][0].strip()
    res = client.post("/api/sessions/owned/messages", json={"content": "hello", "attachment_ids": ["missing"]})
    assert res.status_code == 400


@pytest.fixture
def store(tmp_path):
    from src.chat_attachments.store import AttachmentStore
    return AttachmentStore(tmp_path / "attachments")


def ready(store, session="chat", text="Attachment fact", name="facts.txt"):
    a = store.create(session, name, text.encode())
    store.finish(session, a.id, content=text, position_map=[], images=[], parser="local")
    return store.get(session, a.id)


def test_source_and_parse_survive_reopen(store):
    a = ready(store)
    reopened = type(store)(store.root)
    assert reopened.get("chat", a.id).status == "ready"
    assert reopened.source_path("chat", a.id).read_bytes() == b"Attachment fact"
    assert reopened.text("chat", a.id) == "Attachment fact"


def test_ids_are_session_scoped_and_path_safe(store):
    a = ready(store)
    assert store.get("other", a.id) is None
    with pytest.raises(ValueError):
        store.get("chat", "../escape")
    with pytest.raises(ValueError):
        store.create("../escape", "safe.txt", b"x")
    assert store.create("chat", "../../safe.txt", b"x").filename == "safe.txt"


def test_chat_local_types_do_not_depend_on_mineru_configuration():
    from src.chat_attachments import service
    assert service.capabilities(None) == service.capabilities(SimpleNamespace(enabled=True, api_token="token"))
    assert service.capabilities()["cloud_only"] == []


def test_invalid_pdf_reports_a_local_failure_without_cloud_fallback(store):
    from src.chat_attachments import service
    a = store.create("chat", "invalid.pdf", b"invalid pdf")
    service.parse_attachment(store, "chat", a.id, SimpleNamespace(enabled=True, api_token="token"))
    result = store.get("chat", a.id)
    assert result.status == "failed"
    assert result.error
    assert not result.warnings


def test_text_is_local_and_unsupported_legacy_files_fail_explicitly(store):
    from src.chat_attachments import service
    cfg = SimpleNamespace(enabled=True, api_token="token")
    a = store.create("chat", "notes.txt", b"hello")
    service.parse_attachment(store, "chat", a.id, cfg)
    assert store.text("chat", a.id) == "hello"
    b = store.create("chat", "old.doc", b"old")
    service.parse_attachment(store, "chat", b.id, cfg)
    assert store.get("chat", b.id).status == "failed"
    assert "local" in store.get("chat", b.id).error.lower()


def test_long_attachments_have_searchable_full_text_and_bounded_context(store):
    from src.chat_attachments.service import attachment_context, execute_attachment_tool
    a = ready(store, text="head\n" + "ordinary line\n" * 9000 + "needle: 407\n")
    store.bind("chat", [a.id])
    rendered = attachment_context(store, "chat", [a.id], is_visual=False)
    assert len(json.dumps(rendered)) < 24000
    hits = json.loads(execute_attachment_tool(store, "chat", "search_chat_attachment", {"attachment_id": a.id, "query": "needle"}))
    assert hits["matches"][0]["line"] == 9002
    result = json.loads(execute_attachment_tool(store, "chat", "read_chat_attachment", {"attachment_id": a.id, "start_line": 9002}))
    assert "needle: 407" in result["content"]
    assert "head" not in result["content"]
    assert a.id in json.dumps(attachment_context(store, "chat", [], is_visual=False))


def test_long_single_line_can_be_read_without_losing_the_tail(store):
    from src.chat_attachments.service import execute_attachment_tool
    text = "x" * 25000 + "tail"
    a = ready(store, text=text)
    store.bind("chat", [a.id])
    first = json.loads(execute_attachment_tool(store, "chat", "read_chat_attachment", {"attachment_id": a.id}))
    assert first["has_more"]
    second = json.loads(execute_attachment_tool(store, "chat", "read_chat_attachment", {
        "attachment_id": a.id, "start_line": first["next_start_line"], "start_char": first["next_start_char"]}))
    assert first["content"] + second["content"] == text
    assert not second["has_more"]


def test_many_ocr_images_and_old_files_have_a_shared_context_budget(store):
    from src.chat_attachments.service import attachment_context, execute_attachment_tool
    a = ready(store)
    store.finish("chat", a.id, content="scan", position_map=[], parser="local",
                 images=[{"id": str(i), "url": "data:image/png;base64,eA==", "ocr": "OCR" * 4000} for i in range(30)])
    store.bind("chat", [a.id])
    for i in range(85):
        b = ready(store, name=f"{i}.txt")
        store.bind("chat", [b.id])
    context = attachment_context(store, "chat", [a.id], is_visual=False)
    assert len(json.dumps(context)) < 30000
    listing = json.loads(execute_attachment_tool(store, "chat", "list_chat_attachments", {}))
    assert listing["total"] == 86
    assert len(listing["attachments"]) <= 30


def test_attachment_tools_cannot_read_unsent_or_foreign_files(store):
    from src.chat_attachments.service import execute_attachment_tool
    a = ready(store)
    for sid in ("chat", "other"):
        result = json.loads(execute_attachment_tool(store, sid, "read_chat_attachment", {"attachment_id": a.id}))
        assert "error" in result


def test_images_remain_visual_even_without_ocr_text(store, monkeypatch):
    from src.chat_attachments import service
    buf = io.BytesIO()
    Image.new("RGB", (64, 64), "red").save(buf, format="PNG")
    a = store.create("chat", "photo.png", buf.getvalue())
    monkeypatch.setattr("src.parsers.image_utils._ocr_image", lambda *a, **kw: ("", 0))
    service.parse_attachment(store, "chat", a.id, None)
    assert store.get("chat", a.id).status == "ready"
    store.bind("chat", [a.id])
    visual = service.attachment_context(store, "chat", [a.id], is_visual=True)
    assert "image_url" in json.dumps(visual)
    textual = json.dumps(service.attachment_context(store, "chat", [a.id], is_visual=False))
    assert "image_url" not in textual
    assert "OCR" in textual
    assert "data:image" not in json.dumps(store.parsed("chat", a.id)), "image bytes should be separate from text/position metadata"
    first = store.image_url("chat", a.id, 0)
    assert first == store.image_url("chat", a.id, 0), "reuse byte-stable image snapshots"


def test_document_image_fences_are_projected_for_the_receiving_model(store):
    from src.chat_attachments.service import attachment_context, execute_attachment_tool
    a = ready(store, text="Text\n:::image\nimage_id: figure\nocr_text: uncertain reading\n:::\n")
    store.finish("chat", a.id, content=store.text("chat", a.id), position_map=[], parser="local",
                 images=[{"id": "figure", "url": "data:image/png;base64,eA==", "ocr": "uncertain reading"}])
    store.bind("chat", [a.id])
    visual = json.dumps(attachment_context(store, "chat", [a.id], is_visual=True))
    assert "image_url" in visual
    assert ":::image" not in visual
    nonvisual = json.dumps(attachment_context(store, "chat", [a.id], is_visual=False))
    assert "uncertain" in nonvisual
    assert ":::image" not in nonvisual
    read = execute_attachment_tool(store, "chat", "read_chat_attachment", {"attachment_id": a.id}, is_visual=True)
    assert ":::image" not in read


def test_failed_and_duplicate_ids_cannot_be_bound(store):
    a = store.create("chat", "bad.txt", b"bad")
    with pytest.raises(ValueError):
        store.bind("chat", [a.id])
    a = ready(store)
    with pytest.raises(ValueError):
        store.bind("chat", [a.id, a.id])


def test_delete_session_purges_only_its_attachments(store):
    a, b = ready(store), ready(store, session="other")
    store.purge("chat")
    assert store.get("chat", a.id) is None
    assert store.get("other", b.id).status == "ready"


def test_existing_chats_with_legacy_ids_work_without_attachments(tmp_path):
    from src.chatbox.agent import ChatboxAgent
    from src.db.sessions import SessionStore
    sessions = SessionStore(str(tmp_path / "legacy.db"))
    sessions.create_session(session_id="legacy:chat")
    agent = ChatboxAgent(sessions, SimpleNamespace(supports_vision=False), None)
    sessions.add_message("legacy:chat", "user", "question")
    assert agent._resolve_tools_and_prompt("direct", "legacy:chat", [])
    assert agent._build_messages("legacy:chat", "question", mode="direct", catalog_text="")[-1]["content"] == "question"


@pytest.mark.parametrize("sid", ["chat", "quick_library", "meeting_record", "group_record"])
def test_agent_projects_attachments_and_offers_scoped_tools(tmp_path, store, sid):
    from src.chatbox.agent import ChatboxAgent
    from src.db.sessions import SessionStore
    sessions = SessionStore(str(tmp_path / "sessions.db"))
    sessions.create_session(session_id=sid)
    a = ready(store, session=sid)
    store.bind(sid, [a.id])
    sessions.add_message(sid, "user", "question", metadata={"attachments": [a.model_dump()]})
    agent = ChatboxAgent(sessions, SimpleNamespace(supports_vision=False), None)
    agent._attachments = store
    tools, _, _ = agent._resolve_tools_and_prompt("direct", sid, [])
    assert "read_chat_attachment" in [t["function"]["name"] for t in tools]
    assert agent._is_allowed_tool("search_chat_attachment", "direct", sid)
    messages = agent._build_messages(sid, "question", mode="direct", catalog_text="")
    assert "Attachment fact" in json.dumps(messages)


@pytest.mark.parametrize("sid,mode", [("chat", "agentic"), ("chat", "direct"), ("quick_library", "direct"), ("meeting_record", "direct"), ("group_record", "direct")])
def test_first_stream_turn_has_attachment_tools_and_persists_references(tmp_path, store, sid, mode, monkeypatch):
    from src.chatbox.agent import ChatboxAgent
    from src.db.sessions import SessionStore
    from tests.test_chatbox_agent import _fake_stream
    sessions = SessionStore(str(tmp_path / "stream.db"))
    sessions.create_session(session_id=sid)
    a = ready(store, session=sid)
    llm = MagicMock()
    llm.supports_vision = False
    llm._client.chat.completions.create.return_value = _fake_stream([(None, {"content": "answer"}), ("stop", {})])
    agent = ChatboxAgent(sessions, llm, None)
    agent._attachments = store
    monkeypatch.setattr(agent, "_has_meetings", lambda *a: False)
    monkeypatch.setattr(agent, "_build_catalog_text", lambda *a, **kw: "")
    monkeypatch.setattr("src.chatbox.agent._build_surface_context", lambda *a: "")
    async def run():
        return [e async for e in agent.chat_stream(sid, "read the file", mode=mode, attachment_ids=[a.id])]
    assert asyncio.run(run())[-1]["type"] == "done"
    tools = llm._client.chat.completions.create.call_args_list[0].kwargs["tools"]
    assert "read_chat_attachment" in [t["function"]["name"] for t in tools]
    user = next(m for m in sessions.get_messages(sid) if m.role == "user")
    assert user.metadata["attachments"][0]["id"] == a.id


def test_text_preview_exposes_page_map_only_for_the_native_text_prefix(api):
    from src.chat_attachments.store import for_sessions
    client, sessions = api
    store = for_sessions(sessions)
    a = store.create("owned", "Pages.pdf", b"source")
    points = [{"char_offset": 0, "page_number": 1}, {"char_offset": 9, "page_number": 6}]
    store.finish("owned", a.id, content="Opening.\nStructure.", position_map=points, images=[], parser="local", native=True)
    route = f"/api/sessions/owned/attachments/{a.id}/text"
    assert client.get(route).json()["position_map"] == points
    store.finish("owned", a.id, content="Rewritten by OCR.", position_map=points, images=[], parser="local")
    assert client.get(route).json()["position_map"] == []
