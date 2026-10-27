"""A shared safe document tree for HTML, editable Word and Writer PDF."""
from __future__ import annotations

import base64
import html
import io
import re
import subprocess
import tempfile
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.chatbox.citations import _MARKERS, citation_ids
from src.reports.models import MAX_ASSETS, MAX_ASSET_BYTES, MAX_TOTAL_ASSETS

_STYLE = """
@page{size:A4;margin:20mm}*{box-sizing:border-box}
body{margin:0;background:#f5f7f6;color:#343c37;font:15px/1.85 system-ui,'Noto Sans CJK SC',sans-serif}
article{max-width:960px;margin:32px auto;padding:54px 60px;background:#fff;border-radius:24px}
header{margin-bottom:36px}h1{font-size:30px;line-height:1.4;margin:12px 0;color:#202722;letter-spacing:-.02em}
h2{font-size:21px;line-height:1.5;margin:44px 0 18px;color:#202723}h3{font-size:17px;margin:28px 0 14px;color:#272b28}p{margin:0 0 18px}li{margin:8px 0}
.report-deck{color:#58645d;font-size:15px;line-height:1.8;margin-bottom:28px}.report-deck strong{font-weight:500}.report-subheading{margin:26px 0 12px}.report-caption{margin:10px 0 28px;color:#77817b;font-size:12px;line-height:1.7}.report-caption strong{font-weight:500}
.table-scroll{max-width:100%;overflow-x:auto;margin:24px 0;border:1px solid #eaf0ec;border-radius:14px}
table{width:max-content;min-width:100%;border-collapse:collapse;font-size:13px;line-height:1.6;font-variant-numeric:tabular-nums}
th{background:#f1f3f1;text-align:left;font-weight:600}th,td{padding:10px 12px;max-width:320px;border-bottom:1px solid #e4e7e4;vertical-align:top;word-break:normal;overflow-wrap:normal}tr:last-child td{border-bottom:0}
img{max-width:100%;height:auto}a{color:#1a5e3d}sup a{font-size:11px;text-decoration:none;background:#edf3ef;padding:1px 4px;border-radius:3px}
blockquote{margin:24px 0;padding:18px 22px;border-radius:16px;background:#f5f7f6;color:#58645d}blockquote p:last-child{margin-bottom:0}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f5f6f5;padding:14px;font-size:12px}
.meta,.location{color:#747d77;font-size:12px}.sources{margin-top:44px}.source{margin:26px 0}.source img{display:block;margin:16px 0}.source .location{margin-top:6px}
h1,h2,h3{break-after:avoid}tr{break-inside:avoid}thead{display:table-header-group}
@page wide{size:A4 landscape;margin:18mm}
nav{margin:30px 0;padding:18px 22px;background:#f5f7f6;border-radius:16px;font-size:13px}nav h2{margin:0 0 12px;font-size:13px}nav li{margin:5px 0}nav a{text-decoration:none}
@media print{.table-scroll.is-wide{page:wide;break-before:page;break-after:page}.is-wide table{font-size:10px}body{background:white}article{max-width:none;margin:0;padding:0;border-radius:0}.table-scroll{overflow:visible}table{width:100%;min-width:0}th,td{padding:7px 8px}}
@media(max-width:600px){article{margin:0;padding:28px 22px}}
"""
_ALLOWED = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "strong", "b", "em", "i", "del",
            "blockquote", "pre", "code", "table", "thead", "tbody", "tr", "th", "td", "hr", "br", "a", "sup", "img"}


def location_label(evidence):
    loc = evidence.get("location") or {}
    parts = []
    for key, label in (("page_number", "页"), ("slide_number", "幻灯片"), ("paragraph_index", "段落"), ("line", "行")):
        if loc.get(key) is not None:
            parts.append(f"{label} {loc[key]}")
    if loc.get("sheet_name"):
        parts.append(str(loc["sheet_name"]))
    if loc.get("start_time") is not None:
        def stamp(value):
            seconds = max(0, int(float(value)))
            return f"{seconds // 60:02d}:{seconds % 60:02d}"
        parts.append(stamp(loc["start_time"]) + ("–" + stamp(loc["end_time"]) if loc.get("end_time") is not None else ""))
    return " · ".join(parts)


class _Assets:
    def __init__(self):
        self.urls = set()
        self.total = 0

    def embed(self, data):
        from PIL import Image
        if not data or len(data) > MAX_ASSET_BYTES:
            raise ValueError("Report image asset exceeds the supported limit")
        with Image.open(io.BytesIO(data)) as raster:
            if raster.width * raster.height > 40_000_000:
                raise ValueError("Report image dimensions exceed the supported limit")
            raster.thumbnail((2400, 2400))
            output = io.BytesIO()
            raster.convert("RGBA" if "A" in raster.getbands() else "RGB").save(output, format="PNG")
        data = output.getvalue()
        url = "data:image/png;base64," + base64.b64encode(data).decode()
        if url not in self.urls:
            if len(self.urls) >= MAX_ASSETS or len(data) > MAX_ASSET_BYTES or self.total + len(data) > MAX_TOTAL_ASSETS:
                raise ValueError("Report image assets exceed the supported limit")
            self.urls.add(url)
            self.total += len(data)
        return url


def _safe_tree(source, images):
    soup = BeautifulSoup(markdown.markdown(source, extensions=["tables", "fenced_code", "sane_lists"]), "html.parser")
    for element in list(soup.find_all(True)):
        if element.parent is None:
            continue
        if element.name in {"script", "style", "iframe", "object", "embed", "form", "svg", "math"}:
            element.decompose()
        elif element.name not in _ALLOWED:
            element.unwrap()
        else:
            attributes = {}
            if element.name == "a" and re.match(r"^(?:https?://|mailto:|#)", str(element.get("href", "")), re.I):
                attributes["href"] = element["href"]
            if element.name == "img":
                if element.get("src") not in images:
                    element.decompose()
                    continue
                attributes = {"src": element["src"], "alt": str(element.get("alt", ""))[:200]}
            element.attrs = attributes
    return soup


def source_tree(evidence, source_images, assets):
    from src.parsers.image_utils import iter_image_fence_fields
    quote = evidence.get("quote") or ""

    def image(fields):
        data = source_images(evidence, fields) if source_images else None
        if data:
            return '<img src="' + assets.embed(data) + '" alt="Source image" />'
        return html.escape(fields.get("ocr_text") or fields.get("description") or "图片暂不可用 / Image unavailable")

    def replace(match):
        fields = iter_image_fence_fields(match.group(0) + "\n:::\n")
        return "\n\n" + image(fields[0] if fields else {}) + "\n\n"

    quote = re.sub(r"^[ \t]*:::image[^\n]*\n.*?(?:^[ \t]*:::[ \t]*(?:\n|$)|\Z)", replace, quote.replace("\r\n", "\n"), flags=re.M | re.S)
    if not quote.strip():
        quote = image({}) if (evidence.get("location") or {}).get("image_index") is not None else "页图依据 / Page image evidence"
    return _safe_tree(quote, assets.urls)


def document_tree(report, assets=None, *, image_resolver=None):
    assets = assets or _Assets()
    source = str(report.get("markdown") or "")
    def inline_image(match):
        url = match.group(1)
        if match.group(2) and image_resolver:
            node = BeautifulSoup(match.group(2), 'html.parser').find('img')
            url = str(node.get('src') or '') if node else ''
        if url is None:
            return match.group(0)
        if image_resolver is None:
            raise ValueError("Report image references are unsupported; use self-contained fenced SVG diagrams")
        return '<img src="' + assets.embed(image_resolver(url)) + '" alt="Image" />'

    source = re.sub(r"```.*?(?:```|\Z)|`[^`\n]*`|!\[[^\]\n]*\]\(([^)]+)\)|(<img\b[^>]*>)", inline_image, source, flags=re.S)
    citations = {e["id"]: e for e in report.get("citations", [])}
    ids = citation_ids(source)
    if missing := [eid for eid in ids if eid not in citations]:
        raise ValueError("Unknown citation: " + ", ".join(missing[:5]))
    numbers = {eid: index + 1 for index, eid in enumerate(ids)}

    def diagram(match):
        from src.parsers.svg import _render_diagram, _validated_root, _font_options
        import xml.etree.ElementTree as ET
        root = _validated_root(match.group(1))
        family = _font_options().get("font_family", "Noto Sans CJK SC")
        root.set("font-family", family)
        for element in root.iter():
            # Export runtimes do not ship every font named by a generated SVG.
            # Use the bundled CJK family so text survives offline rasterization.
            if "font-family" in element.attrib:
                element.set("font-family", family)
        image = _render_diagram(ET.tostring(root, encoding="unicode"))
        return '\n<img src="' + assets.embed(image.image_bytes) + '" alt="Diagram" />\n'

    source = re.sub(r"^ {0,3}```svg[^\n]*\n(.*?)^ {0,3}```\s*$", diagram, source, flags=re.M | re.S)
    source = _MARKERS.sub(lambda m: f'<sup><a href="#source-{m.group(1)}">{numbers[m.group(1)]}</a></sup>' if m.group(1) else m.group(0), source)
    soup = _safe_tree(source, assets.urls)
    first = soup.find(True, recursive=False)
    if first and first.name == "h1":
        first.decompose()  # The artifact's title field is the single document title.
    introductory = True
    for node in soup.find_all(True, recursive=False):
        if re.fullmatch("h[1-6]", node.name):
            introductory = False
        if node.name != "p":
            continue
        text = node.get_text().strip()
        if re.match(r"^(?:来源|数据来源|资料来源|Sources?|Data source)\s*[：:]", text, re.I):
            node["class"] = "report-caption"
        elif len(node.contents) == 1 and isinstance(node.contents[0], Tag) and node.contents[0].name in {"strong", "b"}:
            if introductory:
                node["class"] = "report-deck"
            elif len(text) < 80 and not re.search(r"[。！？.!?]$", text):
                node["class"] = "report-subheading"
    headings = soup.find_all(["h1", "h2"])
    if report.get('document_kind') != 'meeting' and len(source) >= 6000 and len(headings) >= 4 and not any(h.get_text().strip().casefold() in {"目录", "contents", "table of contents"} for h in headings[:2]):
        for index, heading in enumerate(headings[:32], 1):
            heading["id"] = f"section-{index}"
    return soup, [citations[eid] for eid in ids]


def render_html(report, *, source_images=None):
    assets = _Assets()
    body, citations = document_tree(report, assets)
    refs = []
    for n, e in enumerate(citations, 1):
        quote = source_tree(e, source_images, assets)
        refs.append(f'<section class="source" id="source-{e["id"]}"><strong>{n}. {html.escape(e["title"])}</strong><div class="location">{html.escape(location_label(e))}</div><blockquote>{quote}</blockquote></section>')
    title = html.escape(report["title"])
    headings = body.select('[id^="section-"]')
    contents = '<nav><h2>目录 / Contents</h2><ol>' + ''.join(f'<li><a href="#{h["id"]}">{html.escape(h.get_text())}</a></li>' for h in headings) + '</ol></nav>' if headings else ''
    language = "zh-CN" if re.search(r"[\u4e00-\u9fff]", report["title"]) else "en"
    output = (f'<!DOCTYPE html><html lang="{language}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>{_STYLE}</style></head><body><article><header><div class="meta">SINKDUCE · REPORT</div><h1>{title}</h1></header>{contents}{body}<section class="sources"><h2>来源与依据 / Sources</h2>{"".join(refs) or "<p>未引用外部资料 / No external sources cited.</p>"}</section></article></body></html>')
    soup = BeautifulSoup(output, "html.parser")
    for table in soup.find_all("table"):
        wide = any(len(row.find_all(["th", "td"], recursive=False)) >= 7 for row in table.find_all("tr"))
        table.wrap(soup.new_tag("div", attrs={"class": "table-scroll is-wide" if wide else "table-scroll"}))
    return str(soup).encode()


def render_docx(report, *, source_images=None, image_resolver=None):
    assets = _Assets()
    body, citations = document_tree(report, assets, image_resolver=image_resolver)
    document = Document()
    # The stock Word theme is blue; bind all accent-driven table/border styles
    # to the same restrained palette as the browser report.
    for part in document.part.package.parts:
        if str(part.partname).endswith("/theme/theme1.xml"):
            theme = part.blob.decode()
            colors = ["1A5E3D", "0F4A2D", "747D77", "E4E7E4", "4F5953", "B6C1BA"]
            for index, color in enumerate(colors, 1):
                theme = re.sub(r'(<a:accent' + str(index) + r'>\s*<a:srgbClr val=")[^"]+', lambda m: m.group(1) + color, theme)
            part._blob = theme.encode()
    # Writer also reads the cached colors on table styles. Resolve those
    # explicitly so Word and PDF use the same palette with any theme importer.
    for element in document.styles.element.iter():
        if element.get(qn("w:themeColor")) == "accent1":
            element.set(qn("w:val") if element.tag == qn("w:color") else qn("w:color"), "1A5E3D")
            for key in ("themeColor", "themeTint", "themeShade"):
                element.attrib.pop(qn("w:" + key), None)
        if element.get(qn("w:themeFill")) == "accent1":
            element.set(qn("w:fill"), "F1F3F1")
            for key in ("themeFill", "themeFillTint", "themeFillShade"):
                element.attrib.pop(qn("w:" + key), None)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("272B28")
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(7)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.font.color.rgb = RGBColor.from_string("1A5E3D")
        style.paragraph_format.keep_with_next = True
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    document.styles["Title"].font.color.rgb = RGBColor.from_string("202722")
    document.styles["Title"].font.size = Pt(26)
    document.styles["Heading 1"].font.color.rgb = RGBColor.from_string("0F4A2D")
    document.styles["Heading 2"].font.color.rgb = RGBColor.from_string("272B28")
    subtitle = document.styles["Subtitle"]
    subtitle.font.name = "Noto Sans CJK SC"
    subtitle.font.size = Pt(9)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor.from_string("747D77")
    subtitle.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        ppr = document.styles[name].element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    document.add_heading(report["title"], 0)
    document.add_paragraph('SinkDuce', style="Subtitle")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def internal_link(paragraph, text, anchor, superscript=False):
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), anchor)
        run = paragraph.add_run(text)
        run.font.superscript = superscript
        run.font.color.rgb = RGBColor.from_string("1A5E3D")
        hyperlink.append(run._r)
        paragraph._p.append(hyperlink)

    headings = body.select('[id^="section-"]')
    if headings:
        document.add_heading("目录 / Contents", 1)
        for heading in headings:
            internal_link(document.add_paragraph(), heading.get_text(), heading["id"].replace("-", "_"))

    def inline(paragraph, node, *, bold=False, italic=False, superscript=False):
        if isinstance(node, NavigableString):
            run = paragraph.add_run(str(node))
            run.bold, run.italic, run.font.superscript = bold, italic, superscript
            return
        if not isinstance(node, Tag):
            return
        if node.name == "br":
            paragraph.add_run().add_break()
        elif node.name == "img":
            data = base64.b64decode(node["src"].split(",", 1)[1])
            from PIL import Image
            with Image.open(io.BytesIO(data)) as picture:
                section = document.sections[-1]
                max_width = (section.page_width - section.left_margin - section.right_margin) / 360000
                max_height = (section.page_height - section.top_margin - section.bottom_margin) / 360000 - 2
                width = min(max_width, picture.width / 96 * 2.54, max_height * picture.width / picture.height)
            paragraph.add_run().add_picture(io.BytesIO(data), width=Cm(width))
        elif node.name == "a" and str(node.get("href", "")).startswith("#source-"):
            internal_link(paragraph, node.get_text(), "source_" + node["href"][8:].replace("-", "_"), superscript)
        else:
            for child in node.children:
                inline(paragraph, child, bold=bold or node.name in {"strong", "b"},
                       italic=italic or node.name in {"em", "i"}, superscript=superscript or node.name == "sup")

    source_landscape = False

    def orient(landscape):
        desired = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        if document.sections[-1].orientation != desired:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            section.orientation = desired
            section.page_width, section.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))

    def block(node):
        if isinstance(node, NavigableString):
            if str(node).strip():
                orient(source_landscape)
                document.add_paragraph(str(node))
            return
        if not isinstance(node, Tag):
            return
        if node.name != "table":
            orient(source_landscape)
        if node.name == "table":
            rows = node.find_all("tr")
            columns = max((len(row.find_all(["th", "td"], recursive=False)) for row in rows), default=0)
            if not columns:
                return
            if columns > 12:
                raise ValueError("Report tables support up to 12 columns; split this table")
            orient(columns >= 7 or source_landscape)
            section = document.sections[-1]
            available = section.page_width - section.left_margin - section.right_margin
            weights, minimums = [], []
            for index in range(columns):
                values = [row.find_all(["th", "td"], recursive=False)[index].get_text(" ", strip=True) for row in rows if len(row.find_all(["th", "td"], recursive=False)) > index]
                def measure(value):
                    return sum(2 if ord(c) > 255 else 1 for c in value)
                longest = max((measure(word) for value in values for word in value.split()), default=3)
                minimums.append(Cm(.45 + min(14, longest) * .19))
                weights.append(max(1, min(28, max((measure(value) for value in values), default=1))))
            spare = max(0, available - sum(minimums))
            widths = [int(minimum + spare * weight / sum(weights)) for minimum, weight in zip(minimums, weights)]
            if sum(widths) > available:
                widths = [int(available * width / sum(widths)) for width in widths]
            table = document.add_table(rows=0, cols=columns)
            table.autofit = False
            for column, width in zip(table.columns, widths):
                column.width = width
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                for cell, width, item in zip(cells, widths, row.find_all(["th", "td"], recursive=False)):
                    cell.width = width
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.25
                    for child in item.children:
                        inline(paragraph, child, bold=item.name == "th")
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                    tcpr = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for edge, size in (("top", "100"), ("bottom", "100"), ("left", "110"), ("right", "110")):
                        margin = OxmlElement("w:" + edge)
                        margin.set(qn("w:w"), size)
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    tcpr.append(margins)
                    if item.name == "th" or row_index % 2 == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "F1F3F1" if item.name == "th" else "F8F9F8")
                        tcpr.append(shade)
                if row.find("th"):
                    props.append(OxmlElement("w:tblHeader"))
        elif node.name in {"ul", "ol"}:
            for item in node.find_all("li", recursive=False):
                paragraph = document.add_paragraph(style="List Bullet" if node.name == "ul" else "List Number")
                for child in item.children:
                    if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                        block(child)
                    else:
                        inline(paragraph, child)
        elif node.name == "hr":
            document.add_paragraph("—")
        else:
            level = int(node.name[1]) if re.fullmatch("h[1-6]", node.name) else None
            paragraph = document.add_heading(level=min(level, 3)) if level else document.add_paragraph()
            if node.get("id", "").startswith("section-"):
                bookmark = OxmlElement("w:bookmarkStart")
                bookmark.set(qn("w:id"), str(int(node["id"].split("-")[-1]) + 1000))
                bookmark.set(qn("w:name"), node["id"].replace("-", "_"))
                paragraph._p.insert(1, bookmark)
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), bookmark.get(qn("w:id")))
                paragraph._p.append(end)
            if node.name == "blockquote":
                paragraph.paragraph_format.left_indent = Cm(.5)
            for child in node.children:
                inline(paragraph, child)
            classes = node.get("class", [])
            if "report-caption" in classes:
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(14)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor.from_string("77817B")
            elif "report-deck" in classes:
                paragraph.paragraph_format.space_after = Pt(16)
                for run in paragraph.runs:
                    run.font.bold = False
                    run.font.color.rgb = RGBColor.from_string("58645D")
            elif "report-subheading" in classes:
                paragraph.paragraph_format.space_before = Pt(14)
                paragraph.paragraph_format.keep_with_next = True

    for node in body.contents:
        block(node)
    orient(False)
    if citations or report.get('document_kind') != 'meeting':
        document.add_heading("来源与依据 / Sources", 1)
    for n, evidence in enumerate(citations, 1):
        excerpt = source_tree(evidence, source_images, assets)
        source_landscape = any(len(row.find_all(["th", "td"], recursive=False)) >= 7 for row in excerpt.find_all("tr"))
        orient(source_landscape)
        paragraph = document.add_paragraph(f'{n}. {evidence["title"]}', style="Heading 3")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(n))
        start.set(qn("w:name"), "source_" + evidence["id"].replace("-", "_"))
        paragraph._p.insert(1, start)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(n))
        paragraph._p.append(end)
        document.add_paragraph(location_label(evidence))
        for node in excerpt.contents:
            block(node)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_pdf(report, *, source_images=None, image_resolver=None):
    from src.chat_attachments.rendering import office_binary
    binary = office_binary()
    if not binary:
        raise ValueError("PDF export requires LibreOffice Writer")
    with tempfile.TemporaryDirectory(prefix="sinkduce-report-") as work:
        directory = Path(work)
        source = directory / "report.docx"
        source.write_bytes(render_docx(report, source_images=source_images, image_resolver=image_resolver))
        result = subprocess.run([binary, f'-env:UserInstallation={(directory / "profile").as_uri()}',
            "--headless", "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(directory), str(source)],
            capture_output=True, timeout=120)
        output = directory / "report.pdf"
        if result.returncode or not output.is_file():
            raise ValueError("PDF export requires a working LibreOffice Writer component")
        return output.read_bytes()
